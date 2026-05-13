"""Analyze an uploaded Excel test case template.

Prints content (via pandas) and formatting (via openpyxl) for every sheet — including
merged cells, column widths, row heights, fonts, fills, borders, alignment, and data
validation rules. The onboarding skill uses this output to build TEST_CASE_CONVENTION.md.

Usage:
    python analyze_excel.py <path-to-xlsx>
"""

import sys

import pandas as pd
from openpyxl import load_workbook


def color_str(c):
    if c is None:
        return None
    if c.type == "rgb":
        return c.rgb
    if c.type == "theme":
        return f"theme:{c.theme}"
    if c.type == "indexed":
        return f"indexed:{c.indexed}"
    return str(c)


def analyze(path):
    # Content (all sheets)
    all_sheets = pd.read_excel(path, sheet_name=None, header=None)
    for sheet_name, df in all_sheets.items():
        print(f"=== Content: {sheet_name} ===")
        print(df.to_string())
        print()

    # Formatting (all sheets)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        print(f"=== Format: {ws.title} ===")
        print(f"Dimensions: {ws.dimensions}")
        merged = [str(r) for r in ws.merged_cells.ranges]
        print(f"Merged ranges: {merged or 'none'}")
        widths = {k: v.width for k, v in ws.column_dimensions.items() if v.width}
        heights = {k: v.height for k, v in ws.row_dimensions.items() if v.height}
        print(f"Column widths: {widths}")
        print(f"Row heights: {heights}")

        # Representative cells: first 8 rows
        for row in ws.iter_rows(min_row=1, max_row=min(8, ws.max_row)):
            for cell in row:
                f, fill, b, a = cell.font, cell.fill, cell.border, cell.alignment
                if cell.value is None and fill.patternType is None:
                    continue
                print(
                    f"  {cell.coordinate}={cell.value!r}",
                    f"font=({f.name},{f.size},b={f.bold},i={f.italic},color={color_str(f.color)})",
                    f"fill={fill.patternType}:{color_str(fill.fgColor)}",
                    f"border=(L={b.left.style},R={b.right.style},T={b.top.style},B={b.bottom.style})",
                    f"align=(h={a.horizontal},v={a.vertical},wrap={a.wrap_text})",
                )

        # Data validation (dropdowns / enums)
        for dv in ws.data_validations.dataValidation:
            print(f"  Validation: type={dv.type} formula={dv.formula1} cells={dv.sqref}")
        print()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python analyze_excel.py <path-to-xlsx>", file=sys.stderr)
        sys.exit(1)
    analyze(sys.argv[1])
