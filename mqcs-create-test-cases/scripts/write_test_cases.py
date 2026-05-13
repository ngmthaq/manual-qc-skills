"""Write test cases to an Excel file matching a TEST_CASE_CONVENTION.

The mqcs-create-test-cases skill parses TEST_CASE_CONVENTION.md plus the approved test cases
into a single JSON spec, then runs this script to produce the final .xlsx.

Usage:
    python write_test_cases.py <spec.json>

Spec JSON schema (all style fields optional — omit to use openpyxl defaults):

{
  "output_path": "/mnt/user-data/outputs/13-5-2026-14-30-15-foo.xlsx",
  "sheets": [
    {
      "name": "Test Cases",
      "header_row": 1,
      "columns": ["ID", "Title", "Steps", "Expected", "Priority", "Status"],
      "column_widths": {"A": 12, "B": 40, "C": 60, "D": 50, "E": 10, "F": 12},
      "row_heights": {"1": 24},
      "header_style": {
        "font": {"name": "Calibri", "size": 11, "bold": true, "italic": false, "color": "FFFFFFFF"},
        "fill": "FF4F81BD",
        "border": "thin",
        "alignment": {"horizontal": "center", "vertical": "center", "wrap_text": true}
      },
      "data_style": {
        "font": {"name": "Calibri", "size": 11},
        "fill": null,
        "border": "thin",
        "alignment": {"horizontal": "left", "vertical": "top", "wrap_text": true}
      },
      "merged_cells": [],
      "data_validations": [
        {"cells": "E2:E1000", "type": "list", "values": ["High", "Medium", "Low"]}
      ],
      "rows": [
        {"ID": "TC-001", "Title": "...", "Steps": "...", "Expected": "...", "Priority": "High", "Status": ""}
      ]
    }
  ]
}
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


def _border(style):
    if not style or style == "none":
        return Border()
    side = Side(style=style)
    return Border(left=side, right=side, top=side, bottom=side)


def _fill(color):
    if not color or color == "none":
        return PatternFill(fill_type=None)
    return PatternFill("solid", start_color=color, end_color=color)


def _font(spec):
    return Font(
        name=spec.get("name", "Calibri"),
        size=spec.get("size", 11),
        bold=spec.get("bold", False),
        italic=spec.get("italic", False),
        color=spec.get("color"),
    )


def _alignment(spec):
    return Alignment(
        horizontal=spec.get("horizontal"),
        vertical=spec.get("vertical"),
        wrap_text=spec.get("wrap_text", False),
    )


def _apply_style(cell, style):
    if not style:
        return
    if "font" in style:
        cell.font = _font(style["font"])
    if "fill" in style:
        cell.fill = _fill(style["fill"])
    if "border" in style:
        cell.border = _border(style["border"])
    if "alignment" in style:
        cell.alignment = _alignment(style["alignment"])


def build_workbook(spec):
    wb = Workbook()
    wb.remove(wb.active)

    for sheet in spec["sheets"]:
        ws = wb.create_sheet(title=sheet["name"])
        header_row = sheet.get("header_row", 1)
        columns = sheet["columns"]

        for col_letter, width in sheet.get("column_widths", {}).items():
            ws.column_dimensions[col_letter].width = width

        for row_num, height in sheet.get("row_heights", {}).items():
            ws.row_dimensions[int(row_num)].height = height

        header_style = sheet.get("header_style", {})
        for col_idx, name in enumerate(columns, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=name)
            _apply_style(cell, header_style)

        data_style = sheet.get("data_style", {})
        for row_offset, row_data in enumerate(sheet.get("rows", [])):
            row_idx = header_row + 1 + row_offset
            for col_idx, name in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(name, ""))
                _apply_style(cell, data_style)

        for rng in sheet.get("merged_cells", []):
            ws.merge_cells(rng)

        for dv_spec in sheet.get("data_validations", []):
            if dv_spec["type"] == "list":
                formula = '"' + ",".join(dv_spec["values"]) + '"'
                dv = DataValidation(type="list", formula1=formula, allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(dv_spec["cells"])

    wb.save(spec["output_path"])
    return spec["output_path"]


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python write_test_cases.py <spec.json>", file=sys.stderr)
        sys.exit(1)
    with open(sys.argv[1]) as f:
        spec = json.load(f)
    path = build_workbook(spec)
    print(f"Wrote: {path}")
